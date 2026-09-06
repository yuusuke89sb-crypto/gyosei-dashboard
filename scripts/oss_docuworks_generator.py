"""
OSS申請 書類一式自動生成エンジン (DocuWorks完全互換PDF / .xdw / 確認書Excel)
行政書士法人フェリス
"""

import os
import sys
import glob
import time
import shutil
import base64
import subprocess
from PIL import Image
import fitz
import openpyxl

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
TEMPLATES_DIR = os.path.join(BASE_DIR, 'templates', 'oss')
DEFAULT_OUTPUT_DIR = os.path.join(BASE_DIR, 'output', 'oss')
FONT_PATH = r"C:\Windows\Fonts\msgothic.ttc"
DW_PRINTER_EXE = r"C:\Program Files (x86)\FUJIFILM\DocuWorks\bin\PDF2DWPrinter.exe"
DW_USER_FOLDER = r"C:\Users\yuusu\Documents\FUJIFILM\DocuWorks\DWFolders\ユーザーフォルダ"

# 16-digit OCR Header Specs:
# Cell width: 19.515 pt, starts at x=69.6 pt
# D1104701 is fixed (cells 0..7)
# cells 8..15 (8 digits) is the order number
CELL_WIDTH = 19.515
CELL_DIGIT_START_X = 225.7

def sanitize_filename(name):
    if not name:
        return '無題'
    for ch in ['\\', '/', ':', '*', '?', '"', '<', '>', '|', '\n', '\r', '\t']:
        name = name.replace(ch, '_')
    return name.strip()

def format_order_no_8(order_no):
    if not order_no:
        return '00000000'
    cleaned = ''.join(c for c in str(order_no) if c.isalnum())
    if len(cleaned) <= 8:
        return cleaned.zfill(8)
    return cleaned[-8:]

def generate_haichi_pdf(data, output_pdf=None):
    master_tpl = os.path.join(TEMPLATES_DIR, 'master_template_haichi.pdf')
    if not os.path.exists(master_tpl):
        raise FileNotFoundError(f"Template not found: {master_tpl}")

    doc = fitz.open(master_tpl)
    page = doc[0]

    order_no_8 = format_order_no_8(data.get('orderNo', ''))
    start_x_h = 225.04
    for i, ch in enumerate(order_no_8):
        cx = start_x_h + i * CELL_WIDTH
        page.insert_text(
            (cx + 5.0, 68.0),
            ch,
            fontname="msgothic",
            fontfile=FONT_PATH,
            fontsize=16.5,
            color=(0, 0, 0)
        )

    # 注文書№ 8マスの縦線（グリッド線）の完全保証描画
    for i in range(9):
        vx = start_x_h + i * CELL_WIDTH
        page.draw_line(fitz.Point(vx, 46.4), fitz.Point(vx, 75.92), color=(0, 0, 0), width=1.44)

    dealer_name = data.get('dealerName', '愛知トヨタ 西春店')
    dealer_tel = data.get('dealerTel', '0568-23-2811')
    dealer_text = f"{dealer_name}  ℡ {dealer_tel}".strip() if dealer_tel else dealer_name
    page.insert_text(
        (575.0, 528.0),
        dealer_text,
        fontname="msgothic",
        fontfile=FONT_PATH,
        fontsize=9.5,
        color=(0, 0, 0)
    )

    reg_no = data.get('regNo', '').strip()
    if reg_no:
        page.insert_text(
            (624.0, 546.0),
            reg_no,
            fontname="msgothic",
            fontfile=FONT_PATH,
            fontsize=11.0,
            color=(0, 0, 0)
        )

    map_src = data.get('haichiMapPath') or data.get('mapImagePath')
    map_b64 = data.get('haichiMapBase64') or data.get('mapImageBase64')

    temp_img_path = None
    if map_b64:
        if ',' in map_b64:
            map_b64 = map_b64.split(',', 1)[1]
        temp_img_path = (output_pdf or os.path.join(DEFAULT_OUTPUT_DIR, 'temp')) + '.temp_map_h.png'
        os.makedirs(os.path.dirname(os.path.abspath(temp_img_path)), exist_ok=True)
        with open(temp_img_path, 'wb') as f:
            f.write(base64.b64decode(map_b64))
        map_src = temp_img_path

    if map_src and os.path.exists(map_src):
        # 配置図描画枠: 原紙枠線の内側に正確に収めて2重線を完全防止
        target_rect = fitz.Rect(71.0, 119.5, 770.5, 513.0)
        page.insert_image(target_rect, filename=map_src, keep_proportion=True)

    if output_pdf:
        os.makedirs(os.path.dirname(os.path.abspath(output_pdf)), exist_ok=True)
        doc.save(output_pdf)

    if temp_img_path and os.path.exists(temp_img_path):
        try: os.remove(temp_img_path)
        except: pass

    return doc

def generate_sozai_pdf(data, output_pdf=None):
    master_tpl = os.path.join(TEMPLATES_DIR, 'master_template_sozai.pdf')
    if not os.path.exists(master_tpl):
        raise FileNotFoundError(f"Template not found: {master_tpl}")

    doc = fitz.open(master_tpl)
    page = doc[0]

    order_no_8 = format_order_no_8(data.get('orderNo', ''))
    start_x_s = 224.40
    for i, ch in enumerate(order_no_8):
        cx = start_x_s + i * CELL_WIDTH
        page.insert_text(
            (cx + 5.0, 72.0),
            ch,
            fontname="msgothic",
            fontfile=FONT_PATH,
            fontsize=16.5,
            color=(0, 0, 0)
        )

    # 所在図 注文書№ 8マスの縦線（グリッド線）の完全保証描画
    for i in range(9):
        vx = start_x_s + i * CELL_WIDTH
        page.draw_line(fitz.Point(vx, 50.6), fitz.Point(vx, 80.84), color=(0, 0, 0), width=1.44)

    map_src = data.get('sozaiMapPath') or data.get('mapImagePath')
    map_b64 = data.get('sozaiMapBase64') or data.get('mapImageBase64')

    temp_img_path = None
    if map_b64:
        if ',' in map_b64:
            map_b64 = map_b64.split(',', 1)[1]
        temp_img_path = (output_pdf or os.path.join(DEFAULT_OUTPUT_DIR, 'temp')) + '.temp_map_s.png'
        os.makedirs(os.path.dirname(os.path.abspath(temp_img_path)), exist_ok=True)
        with open(temp_img_path, 'wb') as f:
            f.write(base64.b64decode(map_b64))
        map_src = temp_img_path

    if map_src and os.path.exists(map_src):
        # 所在図描画枠: 原紙枠線の内側に正確に収めて2重線を完全防止（下部備考欄がないため配置図より縦長）
        target_rect = fitz.Rect(71.0, 123.5, 771.0, 545.5)
        page.insert_image(target_rect, filename=map_src, keep_proportion=True)

    if output_pdf:
        os.makedirs(os.path.dirname(os.path.abspath(output_pdf)), exist_ok=True)
        doc.save(output_pdf)

    if temp_img_path and os.path.exists(temp_img_path):
        try: os.remove(temp_img_path)
        except: pass

    return doc

def convert_pdf_to_xdw(pdf_path, output_xdw):
    if not os.path.exists(DW_PRINTER_EXE):
        print(f"Warning: DocuWorks printer not found at {DW_PRINTER_EXE}")
        return None

    if not os.path.exists(DW_USER_FOLDER):
        os.makedirs(DW_USER_FOLDER, exist_ok=True)

    before_files = set(glob.glob(os.path.join(DW_USER_FOLDER, '*.xdw')))
    
    cmd = [DW_PRINTER_EXE, os.path.abspath(pdf_path)]
    subprocess.run(cmd, capture_output=True, timeout=15)

    for _ in range(10):
        time.sleep(0.5)
        after_files = set(glob.glob(os.path.join(DW_USER_FOLDER, '*.xdw')))
        new_files = after_files - before_files
        if new_files:
            newest_xdw = max(new_files, key=os.path.getmtime)
            os.makedirs(os.path.dirname(os.path.abspath(output_xdw)), exist_ok=True)
            shutil.copy2(newest_xdw, output_xdw)
            return output_xdw

    return None

def generate_combined_pdf(data, output_pdf):
    """
    Page 1: 所在図
    Page 2: 配置図
    """
    doc_sozai = generate_sozai_pdf(data)
    doc_haichi = generate_haichi_pdf(data)

    combined = fitz.open()
    combined.insert_pdf(doc_sozai)
    combined.insert_pdf(doc_haichi)

    os.makedirs(os.path.dirname(os.path.abspath(output_pdf)), exist_ok=True)
    combined.save(output_pdf)
    combined.close()
    doc_sozai.close()
    doc_haichi.close()
    return output_pdf

def generate_confirmation_excel(data, output_xlsx):
    tpl_path = os.path.join(TEMPLATES_DIR, '書類確認書_OSS_原紙.xlsx')
    if not os.path.exists(tpl_path):
        raise FileNotFoundError(f"Excel template not found: {tpl_path}")

    wb = openpyxl.load_workbook(tpl_path)
    ws = wb['編集'] if '編集' in wb.sheetnames else wb.active

    dealer_name = data.get('dealerName', '')
    if dealer_name and not dealer_name.endswith('御中'):
        dealer_display = f"{dealer_name}　　御中"
    else:
        dealer_display = dealer_name

    date_str = data.get('dateStr')
    if not date_str:
        now = time.localtime()
        reiwa_year = now.tm_year - 2018
        date_str = f"令和{reiwa_year}年{now.tm_mon}月{now.tm_mday}日"

    ws.cell(3, 1, dealer_display)
    ws.cell(7, 1, date_str)
    ws.cell(9, 6, data.get('orderNo', ''))
    ws.cell(9, 9, data.get('contactName', ''))
    ws.cell(10, 4, data.get('usageLocation', '申請者に同じ'))
    ws.cell(11, 4, data.get('parkingAddress', '同上'))
    ws.cell(12, 4, data.get('zip', ''))
    ws.cell(13, 4, data.get('carAddress', ''))
    ws.cell(14, 4, data.get('applicantName', ''))
    ws.cell(17, 11, data.get('staffName', ''))

    ws.cell(20, 11, "　　　　　　行政書士法人　フェリス")
    ws.cell(21, 11, "　　　　　　TEL　０５８６－５０－２８９６")
    ws.cell(22, 11, "　　　　　　FAX　０５８６－８７－６６８７")

    os.makedirs(os.path.dirname(os.path.abspath(output_xlsx)), exist_ok=True)
    wb.save(output_xlsx)
    return output_xlsx

def generate_oss_pack(data, output_dir=None):
    """
    一括生成:
    - combined_pdf & combined_xdw (2ページ: 所在図+配置図)
    - haichi_pdf & haichi_xdw
    - sozai_pdf & sozai_xdw
    - excel (書類確認書)
    """
    if not output_dir:
        output_dir = DEFAULT_OUTPUT_DIR

    order_no = data.get('orderNo', 'NoOrder')
    applicant = sanitize_filename(data.get('applicantName', '申請者'))
    prefix = f"【{order_no}】{applicant}"

    res = {
        'combined_pdf': None,
        'combined_xdw': None,
        'haichi_pdf': None,
        'haichi_xdw': None,
        'sozai_pdf': None,
        'sozai_xdw': None,
        'excel': None
    }

    # 1. 2ページ結合 (所在図 + 配置図)
    comb_pdf_path = os.path.join(output_dir, f"{prefix}_所在図配置図.pdf")
    generate_combined_pdf(data, comb_pdf_path)
    res['combined_pdf'] = comb_pdf_path

    comb_xdw_path = os.path.join(output_dir, f"{prefix}_所在図配置図.xdw")
    if convert_pdf_to_xdw(comb_pdf_path, comb_xdw_path):
        res['combined_xdw'] = comb_xdw_path

    # 2. 配置図 個別
    haichi_pdf_path = os.path.join(output_dir, f"{prefix}_配置図.pdf")
    generate_haichi_pdf(data, haichi_pdf_path)
    res['haichi_pdf'] = haichi_pdf_path

    haichi_xdw_path = os.path.join(output_dir, f"{prefix}_配置図.xdw")
    if convert_pdf_to_xdw(haichi_pdf_path, haichi_xdw_path):
        res['haichi_xdw'] = haichi_xdw_path

    # 3. 所在図 個別
    sozai_pdf_path = os.path.join(output_dir, f"{prefix}_所在図.pdf")
    generate_sozai_pdf(data, sozai_pdf_path)
    res['sozai_pdf'] = sozai_pdf_path

    sozai_xdw_path = os.path.join(output_dir, f"{prefix}_所在図.xdw")
    if convert_pdf_to_xdw(sozai_pdf_path, sozai_xdw_path):
        res['sozai_xdw'] = sozai_xdw_path

    # 4. 書類確認書 Excel
    dealer_name = sanitize_filename(data.get('dealerName', '愛知トヨタ'))
    excel_path = os.path.join(output_dir, f"書類確認書（OSS）_{dealer_name}_{order_no}_{applicant}.xlsx")
    generate_confirmation_excel(data, excel_path)
    res['excel'] = excel_path

    return res

if __name__ == '__main__':
    import json
    import argparse

    parser = argparse.ArgumentParser(description="OSS DocuWorks Pack Generator")
    parser.add_argument('--input', '-i', type=str, help="Path to JSON payload file")
    parser.add_argument('--output', '-o', type=str, help="Output directory")
    args = parser.parse_args()

    sample_data = {
        'orderNo': '58280',
        'dealerName': '愛知トヨタ 西春店',
        'dealerTel': '0568-23-2811',
        'contactName': '山田 太郎',
        'applicantName': '黒田 健一',
        'carAddress': '愛知県北名古屋市西春町1-2-3',
        'parkingAddress': '愛知県北名古屋市西春町1-2-5',
        'regNo': '増　　　車',
        'staffName': '田中'
    }

    if args.input and os.path.exists(args.input):
        with open(args.input, 'r', encoding='utf-8') as f:
            data = json.load(f)
    else:
        data = sample_data

    out_dir = args.output or DEFAULT_OUTPUT_DIR
    print(f"[START] Generating OSS DocuWorks Pack for OrderNo: {data.get('orderNo')} -> {out_dir}...")
    results = generate_oss_pack(data, out_dir)
    print("[SUCCESS] Generation Complete! Generated files:")
    for k, v in results.items():
        print(f"  - [{k}]: {v}")
